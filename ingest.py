"""Parse, chunk, and persist documents locally with deterministic chunk IDs."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from config import Settings
from docx import Document as WordDocument
from filelock import FileLock
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


@dataclass(frozen=True)
class IngestResult:
    document_id: str
    name: str
    chunks: int
    duplicate: bool = False


def safe_filename(name: str) -> str:
    # Handle both separators regardless of the server OS; never write uploaded paths.
    return name.replace("\\", "/").rsplit("/", 1)[-1].strip() or "document.txt"


def parse_document(name: str, content: bytes, max_bytes: int) -> list[Document]:
    name = safe_filename(name)
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file format. Upload a PDF, DOCX, or TXT file.")
    if not content:
        raise ValueError("This file is empty.")
    if len(content) > max_bytes:
        raise ValueError(f"This file exceeds the {max_bytes // (1024 * 1024)} MB upload limit.")
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ValueError("Password-protected PDFs are not supported. Upload an unlocked copy.")
            docs = [
                Document(page_content=page.extract_text() or "", metadata={"source": name, "page": i + 1})
                for i, page in enumerate(reader.pages)
            ]
        elif extension == ".docx":
            word = WordDocument(BytesIO(content))
            # Iterating body elements preserves the order of paragraphs and tables.
            from docx.table import Table
            from docx.text.paragraph import Paragraph

            parts = []
            for element in word.element.body:
                if element.tag.endswith("}p"):
                    parts.append(Paragraph(element, word).text)
                elif element.tag.endswith("}tbl"):
                    table = Table(element, word)
                    parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            docs = [Document(page_content="\n".join(parts), metadata={"source": name})]
        else:
            text = content.decode("utf-8-sig")
            docs = [Document(page_content=text, metadata={"source": name})]
    except UnicodeDecodeError as exc:
        raise ValueError("TXT files must use UTF-8 encoding. Save the file as UTF-8 and try again.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this document. Check that it is a valid, unlocked PDF, DOCX, or TXT file.") from exc
    docs = [doc for doc in docs if doc.page_content.strip()]
    if not docs:
        raise ValueError("No readable text found. Scanned PDFs need OCR before uploading.")
    return docs


class DocumentLibrary:
    def __init__(self, settings: Settings, embeddings=None):
        self.settings = settings
        settings.db_dir.mkdir(parents=True, exist_ok=True)
        self.embeddings = embeddings or OllamaEmbeddings(
            model=settings.embedding_model,
            base_url=settings.ollama_url,
            client_kwargs={"timeout": 180, "trust_env": False},
        )
        self.store = Chroma(
            collection_name=settings.collection_name,
            embedding_function=self.embeddings,
            persist_directory=str(settings.db_dir),
            collection_metadata={"hnsw:space": "cosine"},
        )
        self.lock = FileLock(str(settings.db_dir / f"{settings.collection_name}.lock"), timeout=180)

    def list_documents(self) -> list[dict]:
        with self.lock:
            result = self.store.get(include=["metadatas"])
        documents = {}
        for meta in result["metadatas"]:
            document_id = meta["document_id"]
            if document_id not in documents:
                documents[document_id] = {
                    "id": document_id,
                    "name": meta["source"],
                    "bytes": meta["bytes"],
                    "indexed_at": meta["indexed_at"],
                    "chunks": 0,
                }
            documents[document_id]["chunks"] += 1
        return sorted(documents.values(), key=lambda doc: doc["indexed_at"], reverse=True)

    def ingest(self, name: str, content: bytes) -> IngestResult:
        name = safe_filename(name)
        docs = parse_document(name, content, self.settings.max_upload_bytes)
        document_id = sha256(content).hexdigest()
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.settings.chunk_size, chunk_overlap=self.settings.chunk_overlap
        )
        chunks = splitter.split_documents(docs)
        indexed_at = datetime.now(timezone.utc).isoformat()
        for i, chunk in enumerate(chunks):
            chunk.metadata.update({
                "document_id": document_id, "chunk": i, "bytes": len(content), "indexed_at": indexed_at,
            })
        ids = [f"{document_id}:{i}" for i in range(len(chunks))]
        with self.lock:
            existing = self.store.get(where={"document_id": document_id}, include=["metadatas"])
            if existing["ids"]:
                return IngestResult(document_id, existing["metadatas"][0]["source"], len(existing["ids"]), True)
            try:
                # Batches keep Ollama requests small and below Chroma batch limits.
                for start in range(0, len(chunks), 32):
                    self.store.add_documents(chunks[start:start + 32], ids=ids[start:start + 32])
            except Exception:
                # A failed upload must not leave a partially indexed document.
                self.store.delete(where={"document_id": document_id})
                raise
        return IngestResult(document_id, name, len(chunks))

    def delete(self, document_id: str) -> None:
        with self.lock:
            self.store.delete(where={"document_id": document_id})

    def retrieve(self, question: str, *, document_ids: list[str] | None = None, k: int = 3) -> list[Document]:
        if document_ids == []:
            return []
        filters = {"document_id": {"$in": document_ids}} if document_ids else None
        # Prevent concurrent chat sessions from seeing half-finished uploads.
        with self.lock:
            records = self.store.get(where=filters, include=[])
            if not records["ids"]:
                return []
            return self.store.similarity_search(question, k=min(k, len(records["ids"])), filter=filters)


def process_and_store_document(file_path: str, settings: Settings | None = None) -> IngestResult:
    """Compatibility entry point for the requirements' standalone ingestion script."""
    path = Path(file_path)
    settings = settings or Settings.from_env()
    if path.stat().st_size > settings.max_upload_bytes:
        raise ValueError("This file exceeds the 25 MB upload limit.")
    return DocumentLibrary(settings).ingest(path.name, path.read_bytes())


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Index a PDF, DOCX, or TXT into the local document library.")
    parser.add_argument("files", nargs="+", type=Path)
    args = parser.parse_args()
    for file in args.files:
        try:
            result = process_and_store_document(str(file))
            print(f"{'Already indexed' if result.duplicate else 'Indexed'}: {result.name} ({result.chunks} chunks)")
        except Exception as error:
            parser.exit(1, f"Unable to index {file.name}: {error}\n")
