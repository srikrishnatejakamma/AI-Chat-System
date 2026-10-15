"""Offline integration tests: real parsers and Chroma, deterministic embeddings."""

from dataclasses import replace
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from config import Settings
from chat import NOT_FOUND, build_messages, source_records, stream_answer
from docx import Document as WordDocument
from ingest import DocumentLibrary, parse_document, safe_filename
from langchain_core.embeddings import Embeddings
from local_models import ModelStatus, inspect_ollama
from pypdf import PdfWriter
from support import cleanup_test_directory, stop_test_databases


class TestEmbeddings(Embeddings):
    """A tiny, deterministic embedding implementation that never uses a network."""

    def embed_documents(self, texts):
        return [self.embed_query(text) for text in texts]

    def embed_query(self, text):
        terms = ("budget", "leave", "remote", "equipment", "hours")
        return [float(text.lower().count(term)) for term in terms] + [0.1]


class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(cleanup_test_directory, self.temp)
        self.addCleanup(stop_test_databases)
        self.settings = Settings(data_dir=Path(self.temp.name))
        self.library = DocumentLibrary(self.settings, embeddings=TestEmbeddings())

    def test_deduplication_and_persistence(self):
        first = self.library.ingest("policy.txt", b"The learning budget is INR 30,000.")
        duplicate = self.library.ingest("renamed.txt", b"The learning budget is INR 30,000.")
        self.assertFalse(first.duplicate)
        self.assertTrue(duplicate.duplicate)
        self.assertEqual(duplicate.name, "policy.txt")
        reopened = DocumentLibrary(self.settings, embeddings=TestEmbeddings())
        self.assertEqual(len(reopened.list_documents()), 1)
        self.assertTrue((self.settings.db_dir / "chroma.sqlite3").exists())
        self.assertIn("30,000", reopened.retrieve("budget")[0].page_content)

    def test_scoped_retrieval_and_delete(self):
        first = self.library.ingest("budget.txt", b"The annual budget is INR 30,000.")
        second = self.library.ingest("leave.txt", b"The annual leave allowance is 20 days.")
        self.assertEqual(self.library.retrieve("budget", k=1)[0].metadata["document_id"], first.document_id)
        selected = self.library.retrieve("budget", document_ids=[second.document_id])
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0].metadata["source"], "leave.txt")
        self.assertEqual(self.library.retrieve("budget", document_ids=[]), [])
        self.library.delete(first.document_id)
        self.assertEqual(len(self.library.list_documents()), 1)
        self.library.delete(second.document_id)
        self.assertEqual(self.library.retrieve("anything"), [])

    def test_changed_content_with_same_name_is_separate(self):
        self.library.ingest("policy.txt", b"The budget is 100.")
        self.library.ingest("policy.txt", b"The budget is 200.")
        self.assertEqual(len(self.library.list_documents()), 2)

    def test_failed_batch_rolls_back_without_touching_other_documents(self):
        self.library.ingest("good.txt", b"Remote work is available.")
        original_add = self.library.store.add_documents
        calls = 0

        def fail_second_batch(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("Simulated embedding failure")
            return original_add(*args, **kwargs)

        with patch.object(self.library.store, "add_documents", side_effect=fail_second_batch):
            with self.assertRaisesRegex(RuntimeError, "embedding failure"):
                self.library.ingest("large.txt", ("A budget entry for the team. " * 1500).encode())
        self.assertEqual(calls, 2)
        self.assertEqual([doc["name"] for doc in self.library.list_documents()], ["good.txt"])

    def test_embedding_models_use_separate_collections(self):
        self.library.ingest("notes.txt", b"A budget note.")
        other_settings = replace(self.settings, embedding_model="another-local-model")
        other = DocumentLibrary(other_settings, embeddings=TestEmbeddings())
        self.assertEqual(other.list_documents(), [])
        self.assertEqual(len(self.library.list_documents()), 1)

    def test_sources_match_retrieval_numbers(self):
        result = self.library.ingest("notes.txt", b"Annual budget: 30,000.")
        docs = self.library.retrieve("budget")
        sources = source_records(docs)
        self.assertEqual(sources[0]["number"], 1)
        self.assertEqual(sources[0]["document_id"], result.document_id)
        messages = build_messages("How much?", docs, [{"role": "user", "content": "What is the budget?"}])
        self.assertIn("[1] notes.txt", messages[-1].content)
        self.assertIn("untrusted", messages[0].content)


class ParserTests(unittest.TestCase):
    def test_bom_and_path_normalization(self):
        docs = parse_document("../../notes.TXT", b"\xef\xbb\xbfHello", 100)
        self.assertEqual(docs[0].page_content, "Hello")
        self.assertEqual(docs[0].metadata["source"], "notes.TXT")
        self.assertEqual(safe_filename("C:\\fakepath\\notes.txt"), "notes.txt")

    def test_validation(self):
        for name, content, limit, error in [
            ("photo.png", b"x", 100, "Unsupported"),
            ("empty.txt", b"", 100, "empty"),
            ("space.txt", b"  \n ", 100, "No readable text"),
            ("large.txt", b"hello", 2, "exceeds"),
            ("bad.txt", b"\xff\xfe", 100, "UTF-8"),
            ("bad.pdf", b"not a pdf", 100, "Could not read"),
        ]:
            with self.subTest(name=name):
                with self.assertRaisesRegex(ValueError, error):
                    parse_document(name, content, limit)

    def test_docx_paragraphs_and_tables(self):
        word = WordDocument()
        word.add_paragraph("Learning budget")
        table = word.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Amount"
        table.cell(0, 1).text = "30,000"
        word.add_paragraph("Approved annually")
        buffer = BytesIO()
        word.save(buffer)
        doc = parse_document("handbook.docx", buffer.getvalue(), 1024 * 1024)[0]
        self.assertIn("Amount | 30,000", doc.page_content)
        self.assertLess(doc.page_content.index("30,000"), doc.page_content.index("Approved"))

    def test_pdf_page_numbers(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=500, height=500)
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 50 450 Td (The budget is 30000.) Tj ET")
        page[NameObject("/Contents")] = stream
        buffer = BytesIO()
        writer.write(buffer)
        doc = parse_document("test.pdf", buffer.getvalue(), 1024 * 1024)[0]
        self.assertIn("30000", doc.page_content)
        self.assertEqual(doc.metadata["page"], 1)

    def test_scanned_and_encrypted_pdf_errors(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        plain = BytesIO()
        writer.write(plain)
        with self.assertRaisesRegex(ValueError, "OCR"):
            parse_document("scan.pdf", plain.getvalue(), 1024 * 1024)
        writer.encrypt("secret")
        encrypted = BytesIO()
        writer.write(encrypted)
        with self.assertRaisesRegex(ValueError, "Password-protected"):
            parse_document("locked.pdf", encrypted.getvalue(), 1024 * 1024)


class LocalModelTests(unittest.TestCase):
    def test_cloud_endpoints_and_models_rejected(self):
        for kwargs in [{"ollama_url": "https://ollama.com"}, {"chat_model": "gpt-oss:cloud"},
                       {"ollama_url": "http://localhost@evil.example"}, {"ollama_url": "http://localhost/api"}]:
            with self.subTest(kwargs=kwargs):
                with self.assertRaises(ValueError):
                    Settings(**kwargs)

    def test_model_alias_matching(self):
        status = ModelStatus(True, ("nomic-embed-text:latest", "phi3:latest"))
        self.assertTrue(status.has("nomic-embed-text"))
        self.assertFalse(status.has("llama3.1:8b"))

    def test_offline_status(self):
        import httpx
        with patch("local_models.httpx.Client.get", side_effect=httpx.ConnectError("offline")):
            self.assertFalse(inspect_ollama(Settings()).online)

    def test_no_excerpts_never_invokes_model(self):
        with patch("chat.ChatOllama", side_effect=AssertionError("Should not call a model")):
            self.assertEqual("".join(stream_answer(Settings(), "anything", [], [])), NOT_FOUND)


if __name__ == "__main__":
    unittest.main()
